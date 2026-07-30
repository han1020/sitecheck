"""
이전 실행의 엑셀 파일에서 일반점검 행을 가져와,
종료시각이 지나지 않은 것만 다음 실행에 이어 붙이는 모듈.
"""
from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from openpyxl import load_workbook

from .datetime_parser import extract_window

logger = logging.getLogger(__name__)


@dataclass
class CarryoverEntry:
    """이전 엑셀에서 이어온 일반점검 한 줄."""
    kind: str        # '일반점검'
    code: str
    name: str
    schedule: str
    service: str
    reason: str


# excel_writer의 _DATA_START_COL/_HEADER_ROW 와 동기화
_HEADER_ROW = 2
_KIND_COL = 2
_CODE_COL = 3
_NAME_COL = 4
_SCHED_COL = 5
_SVC_COL = 6
_REASON_COL = 7
_SHEET_NAME = "점검"

# 현재 형식: [사이트점검]_YYYYMMDD.xlsx
_FILENAME_RE = re.compile(r"\[사이트점검\]_(\d{8})\.xlsx$")
# 이전 형식(하위호환): maintenance_YYYY-MM-DD_HHMM.xlsx
_LEGACY_RE = re.compile(r"maintenance_(\d{4})-(\d{2})-(\d{2})_(\d{4})\.xlsx$")


def _sort_key(name: str) -> Optional[str]:
    """파일명에서 정렬용 타임스탬프(YYYYMMDDHHMM) 추출. 미인식이면 None.

    macOS 파일시스템이 한글을 NFD(자모 분리) 형태로 저장하는 경우가 있어
    NFC로 정규화한 뒤 정규식 매칭한다.
    """
    name = unicodedata.normalize("NFC", name)
    m = _FILENAME_RE.search(name)
    if m:
        return m.group(1) + "0000"
    m = _LEGACY_RE.search(name)
    if m:
        return m.group(1) + m.group(2) + m.group(3) + m.group(4)
    return None


def find_latest_excel(
    excel_dir: Path,
    exclude_date: Optional[str] = None,
) -> Optional[Path]:
    """디렉토리에서 가장 최근 결과 엑셀 반환 (현재/이전 파일명 형식 모두 인식).

    exclude_date: 'YYYYMMDD' 8자리. 해당 날짜의 엑셀은 후보에서 제외한다.
        (같은 날 재실행 시 오늘 쓴 엑셀을 carryover로 읽어 신규건이 강조를
         잃는 것을 방지하기 위함)
    """
    if not excel_dir.exists():
        return None
    candidates: list[tuple[str, Path]] = []
    for p in excel_dir.glob("*.xlsx"):
        key = _sort_key(p.name)
        if not key:
            continue
        if exclude_date and key[:8] == exclude_date:
            continue
        candidates.append((key, p))
    if not candidates:
        return None
    candidates.sort(key=lambda x: x[0])
    return candidates[-1][1]


def _schedule_end(schedule: str) -> Optional[datetime]:
    """일시 문자열에서 종료 시각 추출. 종료가 없으면 시작 시각 반환."""
    w = extract_window(schedule)
    if not w:
        return None
    return w.end or w.start


def is_active(schedule: str, now: Optional[datetime] = None) -> bool:
    """종료 시각이 now 이후이거나, 시각을 파싱할 수 없으면 활성으로 간주."""
    now = now or datetime.now()
    end = _schedule_end(schedule)
    if end is None:
        # 파싱 실패 = '매월 셋째주 일요일' 같은 반복 표현 또는 형식 미상
        # 일반점검은 보통 구체적 일시이므로, 파싱 실패면 보수적으로 활성 유지
        return True
    return end >= now


def load_previous_general(
    excel_dir: Path,
    now: Optional[datetime] = None,
    exclude_date: Optional[str] = None,
) -> List[CarryoverEntry]:
    """가장 최근 엑셀에서 '일반점검' 행 중 종료 시각이 안 지난 것만 반환.

    exclude_date: 'YYYYMMDD'. 오늘 날짜 엑셀을 carryover 대상에서 제외한다.
    """
    latest = find_latest_excel(excel_dir, exclude_date=exclude_date)
    if latest is None:
        logger.info("이전 엑셀 없음 - carryover 건너뜀")
        return []
    logger.info(f"이전 엑셀 로드: {latest.name}")

    try:
        wb = load_workbook(latest, data_only=True)
    except Exception as e:
        logger.warning(f"이전 엑셀 로드 실패: {e}")
        return []

    if _SHEET_NAME not in wb.sheetnames:
        logger.warning(f"이전 엑셀에 '{_SHEET_NAME}' 시트 없음")
        return []

    ws = wb[_SHEET_NAME]
    carried: List[CarryoverEntry] = []
    dropped = 0

    for row in ws.iter_rows(min_row=_HEADER_ROW + 1, values_only=False):
        kind_cell = row[_KIND_COL - 1] if len(row) >= _KIND_COL else None
        if not kind_cell or kind_cell.value != "일반점검":
            continue

        def _v(col_idx: int) -> str:
            cell = row[col_idx - 1] if len(row) >= col_idx else None
            return str(cell.value).strip() if cell and cell.value is not None else ""

        schedule = _v(_SCHED_COL)
        if not is_active(schedule, now):
            dropped += 1
            continue

        carried.append(
            CarryoverEntry(
                kind="일반점검",
                code=_v(_CODE_COL),
                name=_v(_NAME_COL),
                schedule=schedule,
                service=_v(_SVC_COL),
                reason=_v(_REASON_COL),
            )
        )

    logger.info(f"carryover: {len(carried)}건 유지 / {dropped}건 만료 제거")
    return carried
