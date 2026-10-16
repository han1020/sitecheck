"""Synchronize final Excel rows with the persistence API without overwriting edits."""
from __future__ import annotations

import os
import re
import tempfile
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .datetime_parser import extract_window
from .notice_api import NoticeApi, NoticeApiError
from .notice_identity import (CELL_FIELDS, CHECK_TYPES, identity_from_cells,
                              identity_hash, norm_identity, notice_key)
from .notice_state import NoticeState

KST = timezone(timedelta(hours=9))
ID_COL = 8  # Hidden H column; visible business data remains B:G.
LIMITS = (20, 10, 100, 200, 200, 200)
IMPORT_CHUNK = 1000          # API :import 요청당 최대 행
_NEW_FILL = PatternFill("solid", fgColor="FFFF00")
_NO_FILL = PatternFill(fill_type=None)
_FILE_DATE_RE = re.compile(r"(\d{4})(\d{2})(\d{2})\.xlsx$")


def collection_date_of(path: Path, now: datetime | None = None) -> str:
    """수집일(yyyy-MM-dd, KST). 파일명 [사이트점검]_YYYYMMDD.xlsx 의 날짜, 없으면 오늘."""
    m = _FILE_DATE_RE.search(path.name)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return (now or datetime.now(KST)).date().isoformat()


def exclusion_reason(item: dict[str, Any], now: datetime | None = None) -> str | None:
    """일반점검 중 저장 대상이 아닌 행의 사유. 종료된 점검은 서버가 매일 09:00 물리 삭제하므로
    다시 넣지 않는다. 종료 시각이 없는 일반점검은 서버가 거부하므로 보내지 않는다."""
    if item["checkType"] != "일반점검":
        return None
    if not item["windowEnd"]:
        return "일반점검은 종료 시각이 필요합니다. 일시를 확인하거나 검토 탭에서 다시 입력하세요."
    end = datetime.fromisoformat(item["windowEnd"])
    if end <= (now or datetime.now(KST)):
        return "이미 종료된 점검"
    return None


def configured_sync() -> tuple[NoticeApi, NoticeState] | None:
    api = NoticeApi.from_env()
    if api is None:
        return None
    env = os.environ.get("SITECHECK_API_ENV", "")
    if not env or any(c not in "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-" for c in env):
        raise ValueError("SITECHECK_API_ENV is required and must use letters, digits, _ or -")
    root = Path(__file__).resolve().parent.parent
    return api, NoticeState(root / "output" / "state" / env / "notice_state.json", env)


def item_from_cells(cells: list[Any]) -> dict[str, Any]:
    item = identity_from_cells(cells)
    if item["checkType"] not in CHECK_TYPES:
        raise ValueError("구분은 일반점검 또는 정기점검이어야 합니다.")
    if not all(item[name] for name in ("institutionCode", "institutionName", "scheduleText")):
        raise ValueError("기관코드·기관명·일시는 필수입니다.")
    if any(len(item[name]) > limit for name, limit in zip(CELL_FIELDS, LIMITS)):
        raise ValueError("엑셀 값이 DB 컬럼 길이를 초과했습니다.")
    win = extract_window(item["scheduleText"])
    start = win.start if win else None
    end = win.end if win else None
    if start and end and start > end:
        raise ValueError("시작 시각이 종료 시각보다 늦습니다.")
    for dt in (start, end):
        if dt and dt.microsecond:
            raise ValueError("점검 시각의 소수 초는 저장할 수 없습니다.")
    item["windowStart"] = start.replace(tzinfo=KST).isoformat() if start else None
    item["windowEnd"] = end.replace(tzinfo=KST).isoformat() if end else None
    return item


def _sheet(wb):
    return wb["점검"] if "점검" in wb.sheetnames else wb.worksheets[0]


def _cells(ws, row: int) -> list[str]:
    return ["" if ws.cell(row, col).value is None else str(ws.cell(row, col).value)
            for col in range(2, 8)]


def _write_item(ws, row: int, item: dict[str, Any]) -> None:
    for col, field in enumerate(CELL_FIELDS, 2):
        ws.cell(row, col).value = item[field]
        if "isNew" in item:   # 서버의 신규 표시가 곧 엑셀의 노란색 (응답에 없으면 기존 색 유지)
            ws.cell(row, col).fill = _NEW_FILL if item["isNew"] else _NO_FILL
    ws.cell(row, ID_COL).value = str(item["itemId"])
    ws.column_dimensions["H"].hidden = True


def save_atomic(wb, path: Path) -> None:
    fd, name = tempfile.mkstemp(prefix=".sitecheck-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        wb.save(name)
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def _prepare(api: NoticeApi, state: NoticeState, item: dict[str, Any],
             *, action: str | None, linked_id: str | None) -> dict[str, Any]:
    """저장 전 판정. 보류/판단 필요면 {"status", "candidates"}, 보낼 수 있으면 {"request", ...}."""
    decision = state.decision_for(item)
    from_decision = action is None and decision is not None
    if from_decision:
        action = decision["action"]
        linked_id = decision.get("itemId")
    if action == "HOLD":
        return {"status": "HOLD", "candidates": api.candidates(item)}

    exact = api.lookup([{**{k: item[k] for k in CELL_FIELDS[:5]}, "clientRow": 1}])[0]
    if exact["status"] == "NOT_FOUND" and action not in ("NEW_ALLOW", "LINK"):
        candidates = api.candidates(item)
        if candidates:
            state.decide(item, "HOLD")
            return {"status": "NEEDS_DECISION", "candidates": candidates}
    request = dict(item)
    if action == "LINK":
        if not linked_id:
            raise ValueError("LINK 대상 ID가 없습니다.")
        request["linkedItemId"] = str(linked_id)
    return {"request": request, "decision": decision, "from_decision": from_decision,
            "linked_id": str(linked_id) if action == "LINK" else None}


def resolve_item(api: NoticeApi, state: NoticeState, item: dict[str, Any],
                 *, action: str | None = None, linked_id: str | None = None,
                 retry_on_missing: bool = False) -> dict[str, Any]:
    """한 행을 단건 POST 로 저장한다 (검토 추가·결정 적용용).
    retry_on_missing: 호출자가 넘긴 linked_id 가 엑셀 H열처럼 '기억된 번호'일 때 True.
    사람이 방금 고른 LINK 는 False 로 두어 대상이 없으면 오류를 그대로 올린다."""
    prepared = _prepare(api, state, item, action=action, linked_id=linked_id)
    if "status" in prepared:
        return prepared
    try:
        response = api.create(prepared["request"])
    except NoticeApiError as exc:
        # 연결 대상이 DB에 없음(DB 초기화·물리 삭제로 번호가 바뀜). 기억된 번호만 버리고 처음부터
        # 다시 판정한다 — 식별값 조회는 그대로 하므로 중복 저장은 생기지 않는다.
        # 삭제된 공지는 404 가 아니라 SUPPRESSED_DELETED 로 오므로 이 경로와 겹치지 않는다.
        if prepared["linked_id"] and exc.status == 404 and (prepared["from_decision"] or retry_on_missing):
            if prepared["from_decision"]:
                state.forget(item)
            return resolve_item(api, state, item)
        raise
    if response["status"] == "LINKED_EXISTING" and not prepared["decision"]:
        state.decide(item, "LINK", str(response["item"]["itemId"]))
    return response


def _import_batch(api: NoticeApi, state: NoticeState,
                  batch: list[dict[str, Any]], collection_date: str) -> dict[int, dict[str, Any]]:
    """준비된 행들을 :import 로 보내고 clientRow(엑셀 행) → 결과를 돌려준다.
    연결 대상이 사라진 404 는 그 행의 번호만 버리고 같은 묶음을 다시 보낸다(요청은 전체 롤백이라
    아무것도 저장되지 않은 상태에서 재시도한다)."""
    results: dict[int, dict[str, Any]] = {}
    for offset in range(0, max(len(batch), 1), IMPORT_CHUNK):
        chunk = batch[offset:offset + IMPORT_CHUNK]
        while True:
            try:
                for entry in chunk:
                    entry["request"]["clientRow"] = entry["row"]
                for r in api.import_items([e["request"] for e in chunk], collection_date):
                    results[int(r["clientRow"])] = r
                break
            except NoticeApiError as exc:
                missing = str(exc.details.get("itemId", "")) if exc.status == 404 else ""
                stale = [e for e in chunk if missing and e["linked_id"] == missing]
                if not stale:
                    raise
                for e in stale:
                    e["request"].pop("linkedItemId", None)
                    e["linked_id"] = None
                    if e["from_decision"]:
                        state.forget(e["item"])
                        e["from_decision"] = False
        if not batch:   # 빈 완료 신호는 한 번이면 된다
            break
    return results


def sync_workbook(path: Path, api: NoticeApi, state: NoticeState,
                  collection_date: str | None = None) -> dict[str, Any]:
    """엑셀의 최종 행을 한 번의 :import 로 DB와 맞춘다 (한 수집 실행 = 한 번의 완료 신호).
    보류·판단 필요·종료·검증 실패 행은 보내지 않고 각각 held / skipped / errors 로 보고한다."""
    collection_date = collection_date or collection_date_of(path)
    result: dict[str, Any] = {"saved": 0, "suppressed": 0, "held": [], "errors": [],
                              "skipped": [], "collectionDate": collection_date}
    try:
        api.lookup([])   # 연결·인증 사전 점검: 실패하면 행마다 재시도하며 매달리지 않고 바로 끝낸다
    except NoticeApiError as exc:
        result["errors"].append({"excelRow": 0, "error": f"API 연결 실패: {exc}"})
        result["pendingOperations"] = len(state.pending())
        return result
    now = datetime.now(KST)
    wb = load_workbook(path)
    ws = _sheet(wb)
    batch: list[dict[str, Any]] = []
    try:
        for row in range(3, ws.max_row + 1):
            cells = _cells(ws, row)
            if not any(cells):
                continue
            try:
                item = item_from_cells(cells)
                reason = exclusion_reason(item, now)
                if reason == "이미 종료된 점검":
                    result["skipped"].append({"excelRow": row, "reason": reason})
                    continue
                if reason:
                    raise ValueError(reason)
                existing_id = row_item_id(ws, row)
                decision = state.decision_for(item)
                prepared = _prepare(api, state, item,
                                    action="LINK" if existing_id and not decision else None,
                                    linked_id=existing_id if existing_id and not decision else None)
                if "status" in prepared:
                    result["held"].append({"excelRow": row, "status": prepared["status"],
                                           "candidates": prepared.get("candidates", [])})
                    continue
                batch.append({"row": row, "item": item, **prepared})
            except (ValueError, NoticeApiError, KeyError) as exc:
                result["errors"].append({"excelRow": row, "error": str(exc)})

        try:
            responses = _import_batch(api, state, batch, collection_date)
        except NoticeApiError as exc:
            # 요청 전체가 롤백됐다. 엑셀은 그대로 두고 원인만 보고한다
            result["errors"].append({"excelRow": 0, "error": f"일괄 저장 실패: {exc}"})
            result["pendingOperations"] = len(state.pending())
            return result

        suppressed_rows: list[int] = []
        for entry in batch:
            r = responses.get(entry["row"])
            if r is None:
                result["errors"].append({"excelRow": entry["row"], "error": "응답에 결과가 없습니다."})
                continue
            if r["status"] == "SUPPRESSED_DELETED":
                suppressed_rows.append(entry["row"])
                result["suppressed"] += 1
                continue
            stored = r["item"]
            if r["status"] == "LINKED_EXISTING" and not entry["decision"]:
                state.decide(entry["item"], "LINK", str(stored["itemId"]))
            elif identity_hash(entry["item"]) != identity_hash(stored):
                state.decide(entry["item"], "LINK", str(stored["itemId"]))
            _write_item(ws, entry["row"], stored)
            result["saved"] += 1
        for row in reversed(suppressed_rows):
            ws.delete_rows(row, 1)
        if result["saved"] or suppressed_rows:
            save_atomic(wb, path)
    finally:
        wb.close()
    result["pendingOperations"] = len(state.pending())
    return result

def row_item_id(ws, row: int) -> str | None:
    value = ws.cell(row, ID_COL).value
    return str(value) if value else None


def same_identity(a: dict[str, Any], b: dict[str, Any]) -> bool:
    return identity_hash(a) == identity_hash(b) and all(
        norm_identity(a[field]) == norm_identity(b[field]) for field in CELL_FIELDS[:5])


def append_row(path: Path, cells: list[Any], api: NoticeApi, state: NoticeState,
               *, action: str | None = None, linked_id: str | None = None) -> dict[str, Any]:
    item = item_from_cells(cells)
    resolved = resolve_item(api, state, item, action=action, linked_id=linked_id)
    if resolved["status"] in ("HOLD", "NEEDS_DECISION"):
        return {"ok": False, "needs_decision": True,
                "candidates": resolved.get("candidates", [])}
    if resolved["status"] == "SUPPRESSED_DELETED":
        return {"ok": False, "error": "삭제된 공지는 다시 등록할 수 없습니다."}
    stored = resolved["item"]
    if action in ("LINK", "NEW_ALLOW"):
        state.decide(item, action, linked_id)
    elif identity_hash(item) != identity_hash(stored):
        state.decide(item, "LINK", str(stored["itemId"]))
    wb = load_workbook(path)
    try:
        ws = _sheet(wb)
        for row in range(3, ws.max_row + 1):
            existing_cells = _cells(ws, row)
            if row_item_id(ws, row) == str(stored["itemId"]) or (
                any(existing_cells) and same_identity(identity_from_cells(existing_cells), item)
            ):
                _write_item(ws, row, stored)
                save_atomic(wb, path)
                return {"ok": True, "duplicate": True, "excel_row": row,
                        "item_id": str(stored["itemId"])}
        ws.insert_rows(3)
        _write_item(ws, 3, stored)
        thin = Side(style="thin", color="000000")
        border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for col in range(2, 8):
            cell = ws.cell(3, col)
            cell.font = Font(name="Calibri", size=11)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if col < 5 else "left",
                                       vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="FFFF00")
        ws.row_dimensions[3].height = 32
        save_atomic(wb, path)
    finally:
        wb.close()
    return {"ok": True, "excel_row": 3, "item_id": str(stored["itemId"]),
            "status": resolved["status"]}


def update_row(path: Path, row: int, cells: list[Any], api: NoticeApi,
               state: NoticeState, *, allow_schedule_create: bool = False) -> dict[str, Any]:
    if len(cells) != 6:
        raise ValueError("엑셀의 6개 값을 모두 전달해야 합니다.")
    new_item = item_from_cells(cells)
    wb = load_workbook(path)
    try:
        ws = _sheet(wb)
        if row < 3 or row > ws.max_row:
            raise ValueError("잘못된 행 번호입니다.")
        item_id = row_item_id(ws, row)
        if not item_id:
            raise ValueError("DB 항목 ID가 없습니다. 먼저 해당 엑셀을 API와 동기화하세요.")
        old_item = identity_from_cells(_cells(ws, row))
        changed_schedule = norm_identity(old_item["scheduleText"]) != norm_identity(new_item["scheduleText"])
        if changed_schedule and not allow_schedule_create:
            return {"ok": False, "schedule_change": True,
                    "error": "일시가 달라지면 새 공지로 등록합니다. 기존 공지는 유지됩니다."}
        if changed_schedule:
            operation = {"kind": "CREATE_AFTER_SCHEDULE_CHANGE", "file": path.name,
                         "row": row, "oldItemId": item_id, "payload": new_item}
        else:
            changes = {field: new_item[field] for field in CELL_FIELDS
                       if old_item[field] != new_item[field]}
            if "scheduleText" in changes or "checkType" in changes:
                changes.update(windowStart=new_item["windowStart"],
                               windowEnd=new_item["windowEnd"])
            if not changes:
                return {"ok": True, "updated": 0, "item_id": item_id}
            operation = {"kind": "PATCH", "file": path.name, "row": row,
                         "itemId": item_id, "payload": changes}
        op_id = state.begin(operation)
        try:
            response = api.create(new_item) if changed_schedule else api.patch(item_id, changes)
        except NoticeApiError as exc:
            if exc.status and exc.status < 500:
                state.finish(op_id)
            raise
        if response["status"] == "SUPPRESSED_DELETED":
            state.finish(op_id)
            raise ValueError("삭제된 공지는 다시 등록할 수 없습니다.")
        stored = response["item"]
        if identity_hash(old_item) != identity_hash(stored) and not changed_schedule:
            if notice_key(old_item) == notice_key(stored):
                state.decide(old_item, "LINK", str(stored["itemId"]))
            else:
                state.decide(old_item, "HOLD")
        _write_item(ws, row, stored)
        save_atomic(wb, path)
        state.finish(op_id)
        return {"ok": True, "updated": 1, "item_id": str(stored["itemId"]),
                "new_notice": changed_schedule}
    finally:
        wb.close()


def delete_row(path: Path, row: int, api: NoticeApi, state: NoticeState) -> dict[str, Any]:
    wb = load_workbook(path)
    try:
        ws = _sheet(wb)
        if row < 3 or row > ws.max_row:
            raise ValueError("잘못된 행 번호입니다.")
        item_id = row_item_id(ws, row)
        if not item_id:
            raise ValueError("DB 항목 ID가 없습니다. 먼저 해당 엑셀을 API와 동기화하세요.")
        op_id = state.begin({"kind": "DELETE", "file": path.name, "row": row,
                             "itemId": item_id})
        try:
            api.delete(item_id)
        except NoticeApiError as exc:
            if exc.status and exc.status < 500:
                state.finish(op_id)
            raise
        removed = 0
        for current in range(ws.max_row, 2, -1):
            if row_item_id(ws, current) == item_id:
                ws.delete_rows(current, 1)
                removed += 1
        save_atomic(wb, path)
        state.finish(op_id)
        return {"ok": True, "deleted": removed, "item_id": item_id}
    finally:
        wb.close()


def main() -> None:
    import argparse
    import json

    parser = argparse.ArgumentParser(description="Sync one SiteCheck workbook with the API")
    parser.add_argument("--file", type=Path, help="Workbook to sync; default is the latest")
    parser.add_argument("--date", help="수집일 yyyy-MM-dd (기본: 파일명의 날짜)")
    args = parser.parse_args()
    sync = configured_sync()
    if sync is None:
        parser.error("SITECHECK_API_BASE_URL is required")
    excel_dir = Path(__file__).resolve().parent.parent / "output" / "excel"
    files = sorted(excel_dir.glob("[[]사이트점검[]]_*.xlsx"), reverse=True)
    path = args.file or (files[0] if files else None)
    if path is None or not path.is_file():
        parser.error("No workbook found")
    result = sync_workbook(path, *sync, collection_date=args.date or collection_date_of(path))
    print(json.dumps({"file": str(path), **result}, ensure_ascii=False, indent=2))
    if result["errors"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
