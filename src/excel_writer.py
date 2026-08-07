from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import re

from .carryover import CarryoverEntry
from .config_loader import RegularEntry
from .scraper import NoticeHit

logger = logging.getLogger(__name__)

# 참조 엑셀 시트1의 컬럼/너비 (B2부터 시작)
_COLS = [
    ("구분",    14.5),
    ("기관코드", 13.0),
    ("기관명",  13.0),
    ("일시",    38.0),
    ("업무",    38.3),
    ("사유",    24.6),
]
_LEFT_PAD_COL_WIDTH = 2.88   # A열 너비
_HEADER_ROW = 2
_DATA_START_COL = 2          # B열

_HEADER_FONT = Font(name="Calibri", size=11, bold=True)
_DATA_FONT = Font(name="Calibri", size=11)
_LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="000000")
_BORDER = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)

# 신규(일반)점검 강조 - 노란색
_NEW_FILL = PatternFill("solid", fgColor="FFFF00")


def _kind_for_hit(_hit: NoticeHit) -> str:
    """스크래핑으로 새로 감지된 공지는 '일반점검'으로 표시.
    (정기점검 항목은 RegularEntry에서 별도로 채워짐)"""
    return "일반점검"


def _is_new_hit(h: NoticeHit) -> bool:
    """엑셀 '일반점검(신규)' 행에 들어갈 hit인지.
    정기점검 baseline과 일치한 건(is_regular)은 baseline 행이 이미 있으므로 제외."""
    return bool(h.title) and not getattr(h, "needs_review", False) \
        and not getattr(h, "is_regular", False)


# 제목 키워드 기반 사유 오버라이드.
# 사유는 본문 라벨 추출값(reason_text)을 우선 쓰고, 없으면 제목에서 뽑는다.
# (본문 라벨은 scraper의 _LABELISH_VALUE_RE 가드로 라벨명 누수를 걸러낸다.)
# 둘 다 품질이 안 좋은 공지만 여기서 명시적으로 보정한다. (키워드는 소문자 비교)
REASON_OVERRIDES: list[tuple[str, str]] = [
    # 광주은행 월례 공지 — 제목엔 '온라인 서비스 일시중단'뿐이고 본문에 사유가 있다
    ("광주은행 온라인 서비스 일시중단", "전산시스템 점검 작업"),
    # 홈택스 메인 배너 — 서술 패턴 추출이 '및 손택스 시스템 점검'으로 잘림 (토큰 수 한계)
    ("홈택스 서비스 일시 지연", "홈택스 및 손택스 시스템 점검"),
]


def _resolve_reason_text(title: str, body_reason: str = "") -> str:
    """사유 텍스트 결정. 우선순위: 제목 키워드 오버라이드 → 본문 라벨 추출값 → 제목 정제."""
    t = (title or "").lower()
    for kw, override in REASON_OVERRIDES:
        if kw.lower() in t:
            return override
    if body_reason and body_reason.strip():
        return body_reason.strip()
    return _clean_reason_title(title)


def _clean_reason_title(title: str) -> str:
    """공지 제목에서 사유 컬럼에 들어갈 부분만 추리기.

    제거 항목:
      - 후행 '안내' / '안내드림' / '공지' 등 (예: '시스템 점검 공지' → '시스템 점검')
      - 괄호 안 날짜·일시 (예: '(2026.06.07)', '(6.14.(일) 00:00 ~ 02:00)')
    """
    if not title:
        return ""
    s = title
    # 1) 요일 단독 괄호(예: '(일)', '(토)') 먼저 제거 — 중첩된 외부 괄호 제거를 가능하게 함
    s = re.sub(r"\(\s*[월화수목금토일]\s*\)", "", s)
    # 2) 괄호 안에 숫자(날짜·시간)가 포함된 경우 통째로 제거
    for _ in range(3):
        new = re.sub(r"\s*\([^()]*\d[^()]*\)", "", s)
        if new == s:
            break
        s = new
    # 2-1) 대괄호([...], ［...］, 【...】) 안 내용은 통째로 제거
    #      (예: '[한국투자증권]', '[06/13]', '【공지】')
    for _ in range(3):
        new = re.sub(r"\s*[\[［【][^\[\]［］【】]*[\]］】]\s*", " ", s)
        if new == s:
            break
        s = new
    # 3) 후행 '안내'·'공지' 류 (안내/공지 + 드립니다·드림·합니다 변형, 알림)
    s = re.sub(
        r"\s*(?:안내드립니다|안내드림|안내합니다|공지드립니다|공지드림|공지합니다"
        r"|알림|안내|공지)\s*[.。]?\s*$",
        "",
        s,
    )
    return s.strip()


def _sort_key_for_schedule(schedule_text: str) -> datetime:
    """일반점검 정렬 키: 점검 시작 시각 (파싱 실패 시 datetime.max로 끝에 배치)."""
    from .datetime_parser import extract_window
    w = extract_window(schedule_text) if schedule_text else None
    if w and w.start:
        return w.start
    return datetime.max


def _build_rows(
    hits: List[NoticeHit],
    carried: List[CarryoverEntry],
    regular: List[RegularEntry],
) -> list[tuple]:
    """엑셀에 쓸 행 데이터 리스트 반환.
    튜플 형식: (kind, code, name, schedule, service, reason, is_new, url, shot)
    is_new=True 면 노란색 강조. carryover/정기는 False.

    정렬 정책: '일반점검' 행(carryover + 신규)은 (점검 시작 시각, 기관코드) 오름차순.
              '정기점검' 행은 baseline 순서 그대로 뒤에 붙임.
    """
    rows: list[tuple] = []

    # === 일반점검 항목 통합 ===
    # (sort_key, kind, code, name, schedule, service, reason, is_new, url, shot)
    general_items: list[tuple] = []
    hit_keys: set[tuple[str, str]] = set()
    carried_by_key: dict[tuple[str, str], CarryoverEntry] = {
        (c.code, c.schedule.strip()): c for c in carried
    }

    # 1) 신규 감지(이번 실행에서 잡힌 hit). 본문에서 새로 추출한 값을 사용.
    #    단, 이전 엑셀에 동일 (기관, 일시)가 이미 있었으면 '진짜 신규'가 아니므로
    #    노란색 강조하지 않고 흰색으로 표시한다.
    for h in hits:
        if not _is_new_hit(h):
            continue
        sched = h.schedule_text or ""
        dedup_key = sched.strip() if sched else f"{h.title}|{h.posted_date}"
        hit_keys.add((h.site_code, dedup_key))
        prev = carried_by_key.get((h.site_code, dedup_key))
        service = h.service_text or ""
        reason = _resolve_reason_text(h.title, getattr(h, "reason_text", ""))
        if prev is not None:
            # 이미 이전 엑셀에 있던 행 = 사용자가 업무·사유를 손으로 고쳤을 수 있다.
            # 재수집 값으로 덮어쓰면 수정분이 매번 날아가므로 이전 값을 유지한다.
            # (이전 값이 비어 있을 때만 이번에 추출한 값으로 채운다.)
            service = prev.service or service
            reason = prev.reason or reason
        sort_dt = h.window.start if (h.window and h.window.start) else datetime.max
        general_items.append((
            sort_dt, h.site_code,
            _kind_for_hit(h), h.site_code, h.site_name, sched,
            service, reason,
            prev is None, h.detail_url, h.screenshot_path,
        ))

    # 2) carryover - 신규 hit이 같은 키로 이미 추가했으면 skip (배경 흰색)
    for c in carried:
        if (c.code, c.schedule.strip()) in hit_keys:
            continue
        general_items.append((
            _sort_key_for_schedule(c.schedule), c.code,
            c.kind, c.code, c.name, c.schedule, c.service, c.reason,
            False, "", "",
        ))

    # 일반점검 정렬: (점검 시작 시각, 기관코드) 오름차순
    general_items.sort(key=lambda x: (x[0], x[1]))

    # 정렬 키를 제외한 실제 행 데이터만 rows에 추가
    for it in general_items:
        rows.append(it[2:])

    # 3) 정기점검 baseline (정렬 X, 그대로 추가)
    for r in regular:
        rows.append((
            "정기점검", r.code, r.name, r.schedule, r.service, r.reason,
            False, "", "",
        ))

    return rows


def write_excel(
    hits: List[NoticeHit],
    regular: List[RegularEntry],
    out_path: Path,
    run_id: str,
    carried: Optional[List[CarryoverEntry]] = None,
) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "점검"

    # 좌측 여백 열
    ws.column_dimensions["A"].width = _LEFT_PAD_COL_WIDTH
    # 데이터 컬럼 너비
    for idx, (_, w) in enumerate(_COLS):
        col_letter = get_column_letter(_DATA_START_COL + idx)
        ws.column_dimensions[col_letter].width = w

    # 헤더 (B2~G2)
    for idx, (name, _w) in enumerate(_COLS):
        c = ws.cell(row=_HEADER_ROW, column=_DATA_START_COL + idx, value=name)
        c.font = _HEADER_FONT
        c.alignment = _CENTER
        c.border = _BORDER
    ws.row_dimensions[_HEADER_ROW].height = 22

    # 데이터
    rows = _build_rows(hits, carried or [], regular)
    for r_idx, (kind, code, name, sched, svc, reason, is_new, url, shot) in enumerate(rows):
        row_num = _HEADER_ROW + 1 + r_idx
        values = [kind, code, name, sched, svc, reason]
        for i, v in enumerate(values):
            c = ws.cell(row=row_num, column=_DATA_START_COL + i, value=v)
            c.font = _DATA_FONT
            c.border = _BORDER
            # 구분/기관코드/기관명은 center, 일시/업무/사유는 left
            c.alignment = _CENTER if i < 3 else _LEFT_WRAP
            if is_new:
                c.fill = _NEW_FILL

    # 행 높이 자동 (개략)
    for r_idx in range(len(rows)):
        ws.row_dimensions[_HEADER_ROW + 1 + r_idx].height = 32

    # 헤더 고정
    ws.freeze_panes = ws.cell(row=_HEADER_ROW + 1, column=_DATA_START_COL)

    # === 두 번째 시트: 실행 메타 ===
    meta = wb.create_sheet("실행정보")
    meta["A1"] = "실행 ID"
    meta["B1"] = run_id
    meta["A2"] = "생성 시각"
    meta["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    meta["A3"] = "신규 감지 건수"
    # 이전 엑셀에 이미 있던 건(노란색 없음)은 제외한 '진짜 신규' 수
    meta["B3"] = sum(1 for r in rows if r[6])
    meta["A4"] = "이전 carryover 건수"
    meta["B4"] = len(carried or [])
    meta["A5"] = "에러 사이트 수"
    meta["B5"] = sum(1 for h in hits if h.error)
    meta["A6"] = "정기점검 건수"
    meta["B6"] = len(regular)
    for r in range(1, 7):
        meta.cell(row=r, column=1).font = Font(bold=True)
    meta.column_dimensions["A"].width = 18
    meta.column_dimensions["B"].width = 30

    # 에러가 있으면 상세를 추가
    if any(h.error for h in hits):
        meta["A8"] = "에러 상세"
        meta["A8"].font = Font(bold=True)
        for i, h in enumerate(h for h in hits if h.error):
            meta.cell(row=9 + i, column=1, value=h.site_code)
            meta.cell(row=9 + i, column=2, value=h.site_name)
            meta.cell(row=9 + i, column=3, value=h.error)

    wb.save(out_path)
    logger.info(
        f"엑셀 저장: {out_path} "
        f"(신규 {sum(1 for r in rows if r[6])}건 "
        f"+ 정기 {len(regular)}건)"
    )
    return out_path
