from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

from src.notice_api import NoticeApiError
from src.notice_identity import CELL_FIELDS, identity_hash, notice_key, norm_identity
from src.notice_state import NoticeState
from src.notice_sync import (append_row, delete_row, item_from_cells, sync_workbook,
                             update_row)


CELLS = ["일반점검", "KRBK0101", "저축은행중앙회",
         "2027.03.20(토) 00:00 ~ 06:00", "통합금융정보시스템 전체 업무", "수집 사유"]


class FakeApi:
    """저장 API 흉내. create(단건)·import_items(일괄, collectionDate 필수)·isNew 를 명세대로 따른다."""

    def __init__(self):
        self.items = {}
        self.next_id = 1
        self.import_calls = []

    def lookup(self, items):
        results = []
        for item in items:
            found = next((v for v in self.items.values()
                          if identity_hash(v) == identity_hash(item)), None)
            results.append({"status": "DELETED" if found and found["deletedAt"] else
                            "FOUND" if found else "NOT_FOUND", "item": deepcopy(found)})
        return results

    def candidates(self, item):
        return [deepcopy(v) for v in self.items.values() if notice_key(v) == notice_key(item)]

    def _save_one(self, item, collection_date):
        """단건·일괄 공통 판정. 새 공지는 즉시 저장(같은 요청의 뒤 행은 앞 행 기준으로 판정)."""
        linked = item.get("linkedItemId")
        exact = self.lookup([item])[0]["item"]
        if linked:
            if linked not in self.items:
                raise NoticeApiError("대상 공지를 찾을 수 없습니다.", status=404,
                                     code="ITEM_NOT_FOUND", details={"itemId": linked})
            stored = self.items[linked]
            assert notice_key(stored) == notice_key(item)
            status = "LINKED_EXISTING"
        elif exact:
            stored = exact
            status = "EXISTING" if norm_identity(stored["reasonText"]) == norm_identity(item["reasonText"]) else "REASON_PRESERVED"
        else:
            item_id = str(self.next_id)
            self.next_id += 1
            stored = {key: item[key] for key in CELL_FIELDS}
            stored.update({"itemId": item_id, "windowStart": item["windowStart"],
                           "windowEnd": item["windowEnd"], "deletedAt": None,
                           "isNew": True, "createdDate": collection_date})
            self.items[item_id] = stored
            status = "INSERTED"
        if stored["deletedAt"]:
            status = "SUPPRESSED_DELETED"
        return {"status": status, "item": deepcopy(stored)}

    def create(self, item):
        return self._save_one(item, "2026-01-01")

    def import_items(self, items, collection_date):
        if not collection_date:
            raise NoticeApiError("collectionDate: 필수 값입니다.", status=422, code="VALIDATION_ERROR")
        self.import_calls.append({"n": len(items), "collectionDate": collection_date})
        snapshot, next_id = deepcopy(self.items), self.next_id
        try:
            judged = [(it["clientRow"], self._save_one(it, collection_date)) for it in items]
            for v in self.items.values():          # 이전 날짜의 신규 표시 해제 (같은 트랜잭션)
                if v["isNew"] and v["createdDate"] < collection_date:
                    v["isNew"] = False
            results = []                           # 응답 item 은 해제까지 끝난 최종 저장값
            for client_row, r in judged:
                item = self.items.get(str(r["item"]["itemId"]), r["item"])
                results.append({"clientRow": client_row, "status": r["status"], "item": deepcopy(item)})
            return results
        except Exception:
            self.items, self.next_id = snapshot, next_id   # 전체 롤백
            raise

    def patch(self, item_id, changes):
        self.items[item_id].update(changes)
        return {"status": "UPDATED", "item": deepcopy(self.items[item_id])}

    def delete(self, item_id):
        self.items[item_id]["deletedAt"] = "2026-09-20T00:00:00+09:00"
        return {"status": "DELETED", "item": deepcopy(self.items[item_id])}


def state(tmp_path: Path) -> NoticeState:
    return NoticeState(tmp_path / "state" / "test" / "notice_state.json", "test")


def is_yellow(path: Path, row: int = 3) -> bool:
    wb = load_workbook(path)
    try:
        rgb = wb["점검"].cell(row, 2).fill.fgColor.rgb
        return isinstance(rgb, str) and rgb.upper().endswith("FFFF00")
    finally:
        wb.close()


def sheet_rows(path: Path):
    wb = load_workbook(path)
    try:
        ws = wb["점검"]
        return [[ws.cell(r, c).value for c in range(2, 9)] for r in range(3, ws.max_row + 1)]
    finally:
        wb.close()


def test_repeated_collection_keeps_one_row_and_preserves_reason(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    first = make_xlsx(carried=[CELLS], name="[사이트점검]_20260919.xlsx")
    assert sync_workbook(first, api, local)["saved"] == 1
    api.items["1"]["reasonText"] = "담당자 수정 사유"
    second = make_xlsx(carried=[CELLS], name="[사이트점검]_20260920.xlsx")
    assert sync_workbook(second, api, local)["saved"] == 1
    assert len(api.items) == 1
    assert sheet_rows(second)[0][5:] == ["담당자 수정 사유", "1"]


def test_candidate_needs_explicit_link(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    first = make_xlsx(carried=[CELLS])
    sync_workbook(first, api, local)
    changed = [*CELLS]
    changed[4] = "다르게 수집된 업무"
    second = make_xlsx(carried=[changed], name="[사이트점검]_20260924.xlsx")
    result = sync_workbook(second, api, local)
    assert result["saved"] == 0 and result["held"][0]["candidates"][0]["itemId"] == "1"
    assert local.decision_for(item_from_cells(changed))["action"] == "HOLD"
    local.decide(item_from_cells(changed), "LINK", "1")
    assert sync_workbook(second, api, local)["saved"] == 1
    assert sheet_rows(second)[0][4] == CELLS[4]
    assert len(api.items) == 1


def test_patch_keeps_old_input_link_and_future_collection(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    first = make_xlsx(carried=[CELLS])
    sync_workbook(first, api, local)
    edited = [*CELLS]
    edited[4] = "담당자가 고친 업무"
    assert update_row(first, 3, edited, api, local)["updated"] == 1
    assert local.decision_for(item_from_cells(CELLS))["itemId"] == "1"
    second = make_xlsx(carried=[CELLS], name="[사이트점검]_20260924.xlsx")
    assert sync_workbook(second, api, local)["saved"] == 1
    assert sheet_rows(second)[0][4] == edited[4]


def test_schedule_change_requires_confirmation_and_creates_new(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(carried=[CELLS])
    sync_workbook(path, api, local)
    edited = [*CELLS]
    edited[3] = "2027.03.20(토) 00:00 ~ 07:00"
    assert update_row(path, 3, edited, api, local)["schedule_change"]
    assert len(api.items) == 1
    assert update_row(path, 3, edited, api, local, allow_schedule_create=True)["new_notice"]
    assert len(api.items) == 2
    assert sheet_rows(path)[0][6] == "2"


def test_deleted_item_is_suppressed_on_next_collection(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    first = make_xlsx(carried=[CELLS])
    sync_workbook(first, api, local)
    assert delete_row(first, 3, api, local)["deleted"] == 1
    assert sheet_rows(first) == []
    second = make_xlsx(carried=[CELLS], name="[사이트점검]_20260924.xlsx")
    assert sync_workbook(second, api, local)["suppressed"] == 1
    assert sheet_rows(second) == []
    assert len(api.items) == 1


def test_review_append_uses_identity_not_code_and_reason(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx()
    first = append_row(path, CELLS, api, local)
    assert first["ok"] and not first.get("duplicate")
    second = [*CELLS]
    second[3] = "2027.03.21(일) 00:00 ~ 06:00"
    assert append_row(path, second, api, local)["ok"]
    assert len(api.items) == 2
    assert append_row(path, CELLS, api, local)["duplicate"]
    assert len(sheet_rows(path)) == 2


def test_regular_schedule_without_date_has_null_windows(make_xlsx, tmp_path):
    regular = ["정기점검", "KRBK0101", "저축은행중앙회",
               "매월 셋째 주 일요일 02:00 ~ 04:00", "인터넷뱅킹", "정기 시스템 점검"]
    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(regular=[regular])
    assert sync_workbook(path, api, local)["saved"] == 1
    assert api.items["1"]["windowStart"] is None
    assert api.items["1"]["windowEnd"] is None


def test_web_edit_delete_and_sort_keep_hidden_id(make_xlsx, tmp_path, monkeypatch):
    from src import webserver

    api, local = FakeApi(), state(tmp_path)
    later = [*CELLS]
    later[3] = "2027.03.21(일) 00:00 ~ 06:00"
    path = make_xlsx(carried=[later, CELLS])
    sync_workbook(path, api, local)
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path)
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))
    row_to_id = {row[3]: row[6] for row in sheet_rows(path)}
    assert webserver.sort_excel_rows(path.name)["ok"]
    assert {row[3]: row[6] for row in sheet_rows(path)} == row_to_id
    first_id = sheet_rows(path)[0][6]
    edited = [str(v) for v in sheet_rows(path)[0][:6]]
    edited[5] = "웹에서 수정"
    assert webserver.update_excel_rows(path.name, [{"excel_row": 3, "cells": edited}])["ok"]
    assert api.items[first_id]["reasonText"] == "웹에서 수정"
    assert webserver.delete_excel_row(path.name, 3)["ok"]
    assert api.items[first_id]["deletedAt"]
    assert len(sheet_rows(path)) == 1


def test_web_review_append_requires_decision_for_candidate(make_xlsx, tmp_path, monkeypatch):
    from src import webserver

    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(carried=[CELLS])
    sync_workbook(path, api, local)
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path)
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))
    changed = [*CELLS]
    changed[4] = "검토 추가 업무"
    response = webserver.append_excel_row(path.name, changed)
    assert response["needs_decision"] and len(api.items) == 1
    linked = webserver.append_excel_row(path.name, changed, action="LINK", linked_id="1")
    assert linked["duplicate"]
    assert len(api.items) == 1


def reset_db(api: FakeApi, next_id: int = 100) -> None:
    """DB 초기화 흉내: 모든 행이 사라지고 번호가 새로 시작한다."""
    api.items.clear()
    api.next_id = next_id


def test_stale_hidden_id_after_db_reset_is_replaced(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(carried=[CELLS])
    sync_workbook(path, api, local)
    assert sheet_rows(path)[0][6] == "1"
    reset_db(api)                                       # H열의 1번은 이제 DB에 없다
    result = sync_workbook(path, api, local)
    assert result["errors"] == [] and result["saved"] == 1
    assert sheet_rows(path)[0][6] == "100"              # 새 번호로 교체, 새 행 1개
    assert len(api.items) == 1


def test_stale_link_decision_is_forgotten_after_db_reset(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    sync_workbook(make_xlsx(carried=[CELLS]), api, local)
    changed = [*CELLS]
    changed[4] = "다르게 수집된 업무"
    second = make_xlsx(carried=[changed], name="[사이트점검]_20260924.xlsx")
    sync_workbook(second, api, local)                   # 후보 → HOLD
    local.decide(item_from_cells(changed), "LINK", "1")
    assert sync_workbook(second, api, local)["saved"] == 1
    reset_db(api)                                       # 결정이 가리키는 1번이 사라짐
    fresh = make_xlsx(carried=[changed], name="[사이트점검]_20260925.xlsx")
    result = sync_workbook(fresh, api, local)
    assert result["errors"] == [] and result["saved"] == 1
    assert local.decision_for(item_from_cells(changed)) is None   # 틀린 기억은 지워짐
    assert sheet_rows(fresh)[0][6] == "100"


def test_explicit_link_to_missing_item_is_an_error(make_xlsx, tmp_path):
    import pytest
    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx()
    with pytest.raises(NoticeApiError) as caught:       # 사람이 고른 대상이 없으면 조용히 신규로 바꾸지 않는다
        append_row(path, CELLS, api, local, action="LINK", linked_id="999")
    assert caught.value.status == 404
    assert len(api.items) == 0


class DownApi(FakeApi):
    """API 서버가 죽어 있는 상황: 모든 호출이 연결 오류."""

    def __init__(self):
        super().__init__()
        self.calls = 0

    def lookup(self, items):
        self.calls += 1
        raise NoticeApiError("API connection failed: connection refused")

    def candidates(self, item):
        raise NoticeApiError("API connection failed: connection refused")

    def create(self, item):
        raise NoticeApiError("API connection failed: connection refused")

    def delete(self, item_id):
        raise NoticeApiError("API connection failed: connection refused")


def test_sync_workbook_stops_early_when_api_unreachable(make_xlsx, tmp_path):
    api, local = DownApi(), state(tmp_path)
    later = [*CELLS]
    later[3] = "2027.03.21(일) 00:00 ~ 06:00"
    path = make_xlsx(carried=[CELLS, later])
    result = sync_workbook(path, api, local)
    assert result["saved"] == 0 and len(result["errors"]) == 1
    assert result["errors"][0]["excelRow"] == 0
    assert "API 연결 실패" in result["errors"][0]["error"]
    assert api.calls == 1                                   # 사전 점검 1회뿐, 행별 재시도 없음
    assert all(row[6] is None for row in sheet_rows(path))  # 워크북은 건드리지 않음


def test_collect_keeps_excel_and_run_json_when_api_down(tmp_path, monkeypatch):
    from src import webserver
    from src.scraper import NoticeHit

    api, local = DownApi(), state(tmp_path)
    (tmp_path / "excel").mkdir()
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path / "excel")
    monkeypatch.setattr(webserver, "JSON_DIR", tmp_path / "json")
    monkeypatch.setattr(webserver, "SCREENSHOT_DIR", tmp_path / "shots")
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))

    async def fake_scrape_all(**kwargs):
        return [NoticeHit(site_code=CELLS[1], site_name=CELLS[2], category="BK",
                          title="시스템 점검 안내", posted_date="2026-09-01", detail_url="",
                          screenshot_path="", window=None, schedule_text=CELLS[3],
                          service_text=CELLS[4], reason_text=CELLS[5])]
    monkeypatch.setattr(webserver, "scrape_all", fake_scrape_all)

    webserver._run_in_thread()                               # 웹 '지금 수집'이 부르는 함수
    assert webserver.RUN_STATE["status"] == "done"
    assert "DB 동기화 오류 1건" == webserver.RUN_STATE["error"]
    assert list((tmp_path / "excel").glob("*.xlsx"))          # 엑셀은 생성
    assert list((tmp_path / "json").glob("*.json"))           # 감지목록 이력도 남음


def test_sync_sends_one_import_with_collection_date(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    later = [*CELLS]
    later[3] = "2027.03.21(일) 00:00 ~ 06:00"
    path = make_xlsx(carried=[CELLS, later], name="[사이트점검]_20270301.xlsx")
    result = sync_workbook(path, api, local)
    assert result["saved"] == 2 and result["errors"] == []
    assert api.import_calls == [{"n": 2, "collectionDate": "2027-03-01"}]   # 행별이 아니라 한 번
    assert result["collectionDate"] == "2027-03-01"
    assert sorted(r[6] for r in sheet_rows(path)) == ["1", "2"]


def test_empty_completion_signal_is_sent_when_nothing_to_store(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    sync_workbook(make_xlsx(carried=[CELLS]), api, local)
    changed = [*CELLS]
    changed[4] = "다르게 수집된 업무"
    second = make_xlsx(carried=[changed], name="[사이트점검]_20270302.xlsx")
    api.import_calls.clear()
    result = sync_workbook(second, api, local)              # 후보 → 판단 필요 → 보낼 행 없음
    assert result["held"] and result["saved"] == 0
    assert api.import_calls == [{"n": 0, "collectionDate": "2027-03-02"}]   # 그래도 완료 신호는 보낸다


def test_ended_general_notice_is_not_sent(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    ended = [*CELLS]
    ended[3] = "2024.01.06(토) 00:00 ~ 06:00"
    path = make_xlsx(carried=[ended])
    result = sync_workbook(path, api, local)
    assert result["saved"] == 0 and result["errors"] == []
    assert result["skipped"] == [{"excelRow": 3, "reason": "이미 종료된 점검"}]
    assert api.items == {} and api.import_calls[0]["n"] == 0


def test_general_notice_without_end_time_is_reported_not_sent(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    vague = [*CELLS]
    vague[3] = "추후 공지"                                   # 파싱 불가 → 종료 시각 없음
    path = make_xlsx(carried=[vague, CELLS])
    result = sync_workbook(path, api, local)
    assert result["saved"] == 1 and len(result["errors"]) == 1
    assert "종료 시각" in result["errors"][0]["error"]
    assert api.import_calls[0]["n"] == 1                    # 나머지 행은 정상 저장


def test_is_new_from_server_drives_yellow_fill(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(carried=[CELLS], name="[사이트점검]_20270301.xlsx")   # carried = 흰색으로 작성됨
    assert not is_yellow(path)
    sync_workbook(path, api, local)                         # 서버: INSERTED, isNew=True
    assert is_yellow(path)
    nxt = make_xlsx(carried=[CELLS], name="[사이트점검]_20270302.xlsx")
    sync_workbook(nxt, api, local)                          # 다음 날 수집 성공 → 서버가 isNew 해제
    assert not is_yellow(nxt)


def test_stale_link_in_batch_is_dropped_and_resent(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    later = [*CELLS]
    later[3] = "2027.03.21(일) 00:00 ~ 06:00"
    path = make_xlsx(carried=[CELLS, later])
    sync_workbook(path, api, local)
    reset_db(api)
    api.import_calls.clear()
    result = sync_workbook(path, api, local)
    assert result["errors"] == [] and result["saved"] == 2
    assert len(api.import_calls) >= 2                        # 404 → 번호 버리고 재전송
    assert sorted(r[6] for r in sheet_rows(path)) == ["100", "101"]


def test_edit_of_unsaved_row_writes_file_and_saves_through_sync(make_xlsx, tmp_path, monkeypatch):
    """종료 시각이 없어 저장되지 않은 행을 엑셀뷰에서 고치면 파일에 쓰고 바로 DB에 올린다."""
    from src import webserver

    api, local = FakeApi(), state(tmp_path)
    no_end = [*CELLS]
    no_end[3] = "2027.03.20(토) 07:00"                        # 종료 없음 → 저장 안 됨
    path = make_xlsx(carried=[no_end])
    assert sync_workbook(path, api, local)["errors"]
    assert sheet_rows(path)[0][6] is None
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path)
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))
    fixed = [*no_end]
    fixed[3] = "2027.03.20(토) 07:00 ~ 2027.03.21(일) 24:00"
    res = webserver.update_excel_rows(path.name, [{"excel_row": 3, "cells": fixed}])
    assert res["ok"] and res["sync"]["saved"] == 1 and res["sync"]["errors"] == []
    assert sheet_rows(path)[0][3] == fixed[3] and sheet_rows(path)[0][6] == "1"
    assert api.items["1"]["windowEnd"] == "2027-03-22T00:00:00+09:00"


def test_delete_of_unsaved_row_touches_only_the_file(make_xlsx, tmp_path):
    api, local = FakeApi(), state(tmp_path)
    no_end = [*CELLS]
    no_end[3] = "2027.03.20(토) 07:00"
    path = make_xlsx(carried=[no_end])
    assert delete_row(path, 3, api, local) == {"ok": True, "deleted": 1, "item_id": None}
    assert sheet_rows(path) == [] and api.items == {}


def test_row_that_lost_its_hidden_id_is_matched_by_identity(make_xlsx, tmp_path):
    """H열이 비었어도 DB에 같은 공지가 있으면 그 번호로 수정(PATCH)·삭제한다."""
    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(carried=[CELLS])
    sync_workbook(path, api, local)

    def drop_hidden_id():
        wb = load_workbook(path)
        wb["점검"].cell(3, 8).value = None
        wb.save(path)
        wb.close()

    drop_hidden_id()
    edited = [*CELLS]
    edited[5] = "번호 없이 고친 사유"
    assert update_row(path, 3, edited, api, local)["item_id"] == "1"
    assert api.items["1"]["reasonText"] == "번호 없이 고친 사유"   # 새 행이 아니라 PATCH
    assert sheet_rows(path)[0][6] == "1"
    drop_hidden_id()
    assert delete_row(path, 3, api, local)["deleted"] == 1
    assert api.items["1"]["deletedAt"] and sheet_rows(path) == []


def test_collect_skips_db_sync_when_a_site_failed(tmp_path, monkeypatch):
    from src import webserver
    from src.scraper import NoticeHit

    api, local = FakeApi(), state(tmp_path)
    (tmp_path / "excel").mkdir()
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path / "excel")
    monkeypatch.setattr(webserver, "JSON_DIR", tmp_path / "json")
    monkeypatch.setattr(webserver, "SCREENSHOT_DIR", tmp_path / "shots")
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))

    async def fake_scrape_all(**kwargs):
        return [NoticeHit(site_code=CELLS[1], site_name=CELLS[2], category="BK",
                          title="시스템 점검 안내", posted_date="2026-09-01", detail_url="",
                          screenshot_path="", window=None, schedule_text=CELLS[3],
                          service_text=CELLS[4], reason_text=CELLS[5]),
                NoticeHit(site_code="KRBK0999", site_name="응답 없는 은행", category="BK",
                          title="", posted_date="", detail_url="", screenshot_path="",
                          window=None, error="handler: timeout")]
    monkeypatch.setattr(webserver, "scrape_all", fake_scrape_all)

    webserver._run_in_thread()
    assert webserver.RUN_STATE["status"] == "done"
    assert webserver.RUN_STATE["sync"]["status"] == "skipped"
    assert "건너뜀" in webserver.RUN_STATE["error"]
    assert api.import_calls == [] and api.items == {}        # DB는 그대로
    assert list((tmp_path / "excel").glob("*.xlsx"))          # 엑셀·이력은 남는다


def _run_json_with_one_hit(tmp_path: Path, run_id: str) -> Path:
    """감지 목록 run JSON 흉내 (CELLS 공지 한 건) + 삭제 목록용 config 디렉터리."""
    import json
    for name in ("json", "config"):
        (tmp_path / name).mkdir(exist_ok=True)
    hit = {"site_code": CELLS[1], "site_name": CELLS[2], "title": "시스템 점검 안내",
           "schedule_text": CELLS[3], "screenshot_url": ""}
    run_json = tmp_path / "json" / f"{run_id}.json"
    run_json.write_text(json.dumps({"matched": [hit], "count": 1}, ensure_ascii=False),
                        encoding="utf-8")
    return run_json


def test_matched_list_delete_also_soft_deletes_in_db(make_xlsx, tmp_path, monkeypatch):
    import json
    from src import webserver

    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(carried=[CELLS])                        # [사이트점검]_20260923.xlsx
    sync_workbook(path, api, local)
    run_json = _run_json_with_one_hit(tmp_path, "2026-09-23_1000")
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path)
    monkeypatch.setattr(webserver, "JSON_DIR", tmp_path / "json")
    monkeypatch.setattr(webserver, "CONFIG_DIR", tmp_path / "config")
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))

    res = webserver.delete_matched_hit("2026-09-23_1000", 0)
    assert res["ok"] and res["excel_deleted"] == 1 and res["db_deleted"] == 1
    assert api.items["1"]["deletedAt"]                       # 엑셀뷰 삭제와 같은 소프트 삭제
    assert sheet_rows(path) == []
    assert json.loads(run_json.read_text(encoding="utf-8"))["matched"] == []


def test_matched_list_delete_changes_nothing_when_api_down(make_xlsx, tmp_path, monkeypatch):
    import json
    from src import webserver

    api, local = FakeApi(), state(tmp_path)
    path = make_xlsx(carried=[CELLS])
    sync_workbook(path, api, local)
    run_json = _run_json_with_one_hit(tmp_path, "2026-09-23_1000")
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path)
    monkeypatch.setattr(webserver, "JSON_DIR", tmp_path / "json")
    monkeypatch.setattr(webserver, "CONFIG_DIR", tmp_path / "config")
    monkeypatch.setattr(webserver, "configured_sync", lambda: (DownApi(), local))

    res = webserver.delete_matched_hit("2026-09-23_1000", 0)
    assert not res["ok"] and "DB 삭제 실패" in res["error"]
    assert len(sheet_rows(path)) == 1 and not api.items["1"]["deletedAt"]   # 엑셀·DB 그대로
    assert len(json.loads(run_json.read_text(encoding="utf-8"))["matched"]) == 1   # 목록도 그대로
    assert not (tmp_path / "config" / "matched_deletes.yaml").exists()           # 제외 목록도 그대로


def test_excel_view_sync_updates_status_line_of_last_run(tmp_path, monkeypatch):
    """수집 뒤 오류 행을 엑셀뷰에서 고쳐 저장(자동 동기화)하면 상태줄의 DB 결과가 그 결과로 바뀐다."""
    from src import webserver
    from src.datetime_parser import extract_window
    from src.scraper import NoticeHit

    api, local = FakeApi(), state(tmp_path)
    (tmp_path / "excel").mkdir()
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path / "excel")
    monkeypatch.setattr(webserver, "JSON_DIR", tmp_path / "json")
    monkeypatch.setattr(webserver, "SCREENSHOT_DIR", tmp_path / "shots")
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))
    start_only = "2027.03.20(토) 07:00"

    async def fake_scrape_all(**kwargs):
        return [NoticeHit(site_code=CELLS[1], site_name=CELLS[2], category="BK",
                          title="시스템 점검 안내", posted_date="2026-09-01", detail_url="",
                          screenshot_path="", window=extract_window(start_only),
                          schedule_text=start_only, service_text=CELLS[4], reason_text=CELLS[5])]
    monkeypatch.setattr(webserver, "scrape_all", fake_scrape_all)

    webserver._run_in_thread()
    assert webserver.RUN_STATE["error"] == "DB 동기화 오류 1건"          # 종료 시각 없음
    assert webserver.RUN_STATE["sync_at"] == ""
    excel = webserver.RUN_STATE["excel"]
    fixed = [*CELLS]
    fixed[3] = "2027.03.20(토) 07:00 ~ 2027.03.21(일) 24:00"
    assert webserver.update_excel_rows(excel, [{"excel_row": 3, "cells": fixed}])["ok"]
    assert webserver.RUN_STATE["error"] == "" and webserver.RUN_STATE["sync_at"]
    assert webserver.RUN_STATE["sync"]["errors"] == []
    assert webserver.RUN_STATE["sync"]["saved"] >= 1


def test_excel_view_sync_of_another_file_leaves_status_line(make_xlsx, tmp_path, monkeypatch):
    from src import webserver

    api, local = FakeApi(), state(tmp_path)
    other = make_xlsx(carried=[CELLS], name="[사이트점검]_20260101.xlsx")
    monkeypatch.setattr(webserver, "EXCEL_DIR", tmp_path)
    monkeypatch.setattr(webserver, "configured_sync", lambda: (api, local))
    before = {"status": "done", "excel": "[사이트점검]_20260923.xlsx",
              "sync": {"saved": 3, "held": [], "errors": [{"excelRow": 5, "error": "x"}]},
              "sync_at": "", "error": "DB 동기화 오류 1건"}
    monkeypatch.setattr(webserver, "RUN_STATE", {**webserver.RUN_STATE, **before})
    assert webserver.sync_excel_file(other.name)["ok"]
    assert {k: webserver.RUN_STATE[k] for k in before} == before     # 다른 날짜 파일 → 상태줄 그대로
